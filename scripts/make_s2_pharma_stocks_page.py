"""Build the S² Capital interactive pharma stock table (English page) from the Excel workbook.
Usage: python3 scripts/make_s2_pharma_stocks_page.py <workbook.xlsx> public/documents/s2/2026-10-pharma-listed-stocks.html
"""
import json, os, re, sys
from openpyxl import load_workbook
SRC, OUT = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(SRC)
MK = [('a', 'A股', 'A-share', 'CNY'), ('hk', '港股', 'Hong Kong', 'HKD'), ('us', '美股', 'United States', 'USD'), ('ch', '瑞士', 'Switzerland', 'CHF')]
VAL = {
    '化学制药': 'Chemical pharmaceuticals', '医疗器械': 'Medical devices', '中药': 'Traditional Chinese medicine', '生物制品': 'Biologics',
    '医疗服务': 'Healthcare services', '医药商业': 'Pharma distribution & retail', '医疗美容': 'Medical aesthetics',
    '上交所主板': 'SSE Main Board', '上交所科创板': 'SSE STAR Market', '深交所主板': 'SZSE Main Board', '深交所创业板': 'SZSE ChiNext',
    '北京证券交易所': 'Beijing Stock Exchange', '深交所风险警示板': 'SZSE Risk Warning Board', '上交所风险警示板': 'SSE Risk Warning Board',
    '大型制药': 'Pharmaceuticals: Major', '其他制药': 'Pharmaceuticals: Other', '仿制药': 'Pharmaceuticals: Generic', '生物技术': 'Biotechnology',
    '医院/护理管理': 'Hospital/Nursing Management', '管理式医疗': 'Managed Health Care', '医疗/护理服务': 'Medical/Nursing Services',
    '医疗健康行业服务': 'Services to the Health Industry', '医药分销': 'Medical Distributors', '连锁药店': 'Drugstore Chains',
    '补充（东财：药品及生物科技）': 'Added (Eastmoney: Pharma & biotech)', '补充（东财：其他医疗保健）': 'Added (Eastmoney: Other healthcare)', '补充（Nasdaq分类）': 'Added (Nasdaq Health Care)',
    '普通股': 'Common', '存托凭证(ADR)': 'ADR',
    '申万医药生物成分': 'SWS Pharma & Biotech index', '东财行业分类补充': 'Added: Eastmoney classification', 'TradingView行业分类': 'TradingView classification', 'Nasdaq行业分类补充': 'Added: Nasdaq classification',
}
INTL_IND = {'医疗器械': 'Medical Specialties'}  # TradingView's own label on the non-A sheets
def note_en(s):
    s = s.replace('ISIN非瑞士，按SIX为主要上市地纳入', 'Non-Swiss ISIN; included because SIX is its main listing')
    s = s.replace('股价两数据源不一致：', 'Price differs between sources: ').replace('腾讯', 'Tencent')
    s = s.replace('另有交易代码/股份类别：', 'Other ticker / share class: ').replace('Nasdaq行业：', 'Nasdaq industry: ')
    return s.replace('；', '; ').replace('、', ', ')
def head_en(h, cur, a_share):
    per = 'H1 2026' if a_share else 'TTM'
    exact = {'序号': '#', '代码': 'Code', 'ISIN': 'ISIN', '注册国家/地区': 'Country', '行业': 'Industry', '二级行业': 'Industry', '上市板块': 'Board', '交易所': 'Exchange',
             '证券类型': 'Type', '币种': 'Currency', 'PE(TTM)': 'P/E (TTM)', 'PE(动态)': 'P/E (annualised)', 'PE(静态)': 'P/E (last FY)', 'PB': 'P/B', '换手率': 'Turnover',
             '营收同比': 'Revenue YoY', '净利润同比': 'Net profit YoY', '半年报披露日': 'H1 report date', '最近财报期 截止日': 'Latest period end', '备注': 'Notes', '名单来源': 'Source list',
             '港股代码': 'HK code', '港股简称': 'HK name (CN)', '美股代码': 'US ticker', 'B股代码': 'B-share code', 'A股代码': 'A-share code', '瑞士代码': 'SIX ticker',
             '公司中文全称': 'Legal name (CN)', '公司英文全称': 'Company', '公司英文名称': 'Company', '名称': 'Chinese name', '申万医药生物 指数权重': 'SWS index weight'}
    if h in exact: return exact[h]
    if h.startswith('最新价'): return f'Price ({cur})'
    if h.startswith('总市值'): return f'Market cap ({cur} bn)'
    if h.startswith('流通市值'): return f'Free-float cap ({cur} bn)'
    if h.startswith('营业收入'): return f'Revenue {per} ({cur} bn)'
    if h.startswith('归母净利润'): return f'Net profit {per} ({cur} bn)'
    if h.startswith('毛利率'): return 'Gross margin' if a_share else 'Gross margin TTM'
    if h.startswith('ROE'): return f'ROE {per}'
    if h.startswith('每股收益'): return f'EPS {per} ({cur})'
    if h.startswith('每股净资产'): return f'BVPS ({cur})'
    raise SystemExit('untranslated header: ' + h)
CROSS = ('港股代码', '港股简称', '美股代码', 'B股代码', 'A股代码', '瑞士代码')
CATS = ('Industry', 'Board', 'Country', 'Exchange', 'Type', 'Currency', 'Source list')
cjk = lambda s: isinstance(s, str) and any('一' <= ch <= '鿿' for ch in s)
markets = []
for mid, sheet, en, cur in MK:
    ws = wb[sheet]; A = mid == 'a'
    heads = [str(c.value).replace('\n', ' ') for c in ws[1]]
    fmts = [c.number_format for c in ws[2]]
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    eng = heads.index('公司英文全称' if A else '公司英文名称')
    # order: #, code, English company, other identity columns, the rest, then other-market codes last
    order = [0, 1, eng] + [j for j in range(2, len(heads)) if j != eng and heads[j] not in CROSS] + [j for j in range(len(heads)) if heads[j] in CROSS]
    cols, outrows = [], [[] for _ in rows]
    for j in order:
        h = heads[j]; vals = [r[j] for r in rows]
        if h == '名称' and not A: vals = [v if cjk(v) else None for v in vals]  # Chinese name only where one exists
        nn = [v for v in vals if v is not None]
        if not nn: continue
        f = fmts[j] or ''
        isnum = all(isinstance(v, (int, float)) for v in nn) and f != '@'
        t, d = 's', 0
        if isnum and '%' in f: t, d = 'p', (len(f.split('.')[1].rstrip('%')) if '.' in f else 0)
        elif isnum: t, d = 'n', (len(f.split('.')[1]) if '.' in f else 0)
        if h == '序号': t, d = 'n', 0
        if '亿' in h:  # 亿 (100M) -> billions
            vals = [None if v is None else v / 10 for v in vals]; d = 2 if '市值' in h else 3
        if t == 's':
            if h == '备注': vals = [None if v is None else note_en(v) for v in vals]
            elif h in ('行业',) and not A: vals = [INTL_IND.get(v, VAL.get(v, v)) for v in vals]
            else: vals = [VAL.get(v, v) for v in vals]
            if h in ('行业', '二级行业', '上市板块', '证券类型', '名单来源', '备注'):
                bad = sorted({v for v in vals if cjk(v)})
                if bad: raise SystemExit(f'untranslated values in {sheet}/{h}: {bad[:5]}')
        label = head_en(h, cur, A)
        cols.append({'l': label, 't': t, 'd': d, 'c': t == 's' and label in CATS and len(set(nn)) > 1})
        for i, v in enumerate(vals): outrows[i].append(round(v, 6) if isinstance(v, float) else v)
    markets.append({'id': mid, 'label': en, 'cur': cur, 'cols': cols, 'rows': outrows})
    print(en, len(outrows), [c['l'] for c in cols])
n = {m['id']: len(m['rows']) for m in markets}
notes = [
 ['Overview', f"Four market tables: A-share {n['a']:,}, Hong Kong {n['hk']:,}, United States {n['us']:,}, Switzerland {n['ch']:,}. All amounts are in the listing currency; nothing is converted to a common currency. Tables are sorted by market cap by default."],
 ['A-share: universe', '480 stocks are the full constituent list of the Shenwan (SWS) level-1 "Pharma & Biotech" index (801150), with industries taken from the six SWS level-2 indices (801151–801156). A further 36 stocks (29 on the Beijing Stock Exchange, 7 in Shanghai/Shenzhen) are not SWS constituents and were added from Eastmoney\'s industry classification (chemical pharma, TCM, biologics, medical devices, healthcare services, pharma distribution, medical aesthetics); see the "Source list" column. B-shares and NEEQ stocks are excluded.'],
 ['A-share: prices and financials', 'Prices are the 12 Oct 2026 close. Financials are from the 2026 half-year reports (period end 30 Jun 2026): revenue and net profit attributable to shareholders are first-half cumulative, and YoY compares with the first half of 2025. A negative P/E means a loss on that basis. TTM = last four quarters; annualised = latest reporting period annualised; last FY = last full financial year.'],
 ['A-share: cross-listing codes', 'Company names, HK codes and B-share codes come from Eastmoney company profiles. The 23 HK codes were cross-checked against Tencent\'s A+H list (209 pairs). A US ticker is shown only for a US exchange listing, found by matching the English legal name against the SEC ticker list: BeOne Medicines (Nasdaq: ONC) is the only one. OTC ADRs are not shown.'],
 ['Hong Kong, US, Switzerland: universe', 'Stocks whose TradingView sector is Health Technology or Health Services, or whose industry is Medical Distributors or Drugstore Chains. Hong Kong: HKEX Main Board and GEM, HKD counters. United States: common shares and ADRs on NASDAQ, NYSE and AMEX; no preferred shares, no OTC. Switzerland: SIX-listed companies with a Swiss (CH) ISIN, plus three companies with a non-Swiss ISIN whose main listing is SIX (Cosmo, Newron, SHL Telemedicine; flagged in Notes). Secondary SIX listings of foreign companies (e.g. Johnson & Johnson, Pfizer, Sanofi) are excluded.'],
 ['Hong Kong, US: added stocks', 'To reduce gaps from relying on a single classification, 53 Hong Kong stocks were added from Eastmoney\'s HK industries "Pharma & biotech" and "Other healthcare", and 75 US stocks from the Nasdaq Stock Screener\'s Health Care sector (excluding the Nasdaq industries Medicinal Chemicals and Botanical Products, Industrial Specialties and Precision Instruments). Their "Industry" starts with "Added" and the "Source list" column marks them. They include some companies that are not typical pharma (cosmetics, health-tech platforms) and can be filtered out.'],
 ['Hong Kong, US, Switzerland: prices', 'Data retrieved 12 Oct 2026, 11:34 CEST. Hong Kong: 12 Oct 2026 close. United States: 9 Oct 2026 close. Switzerland: intraday price on 12 Oct 2026 (SIX had not closed at retrieval). Hong Kong and US prices were checked stock by stock against Tencent quotes; the six that differ show both values in Notes.'],
 ['Hong Kong, US, Switzerland: financials and valuation', 'Revenue, net profit, gross margin, ROE and EPS are trailing twelve months (TTM); YoY is TTM on TTM. "Latest period end" is the end date of the latest reporting period used by the source. TradingView converts financials into the listing currency (e.g. RMB reporters listed in Hong Kong are shown in HKD). Per-share figures for ADRs are per ADR. A blank P/E means a loss or no data. For dual-listed companies (e.g. H-shares) market cap is the whole company, not only the part traded on that market, so market caps cannot be added across tables. Free-float cap = market cap × free-float percentage.'],
 ['Hong Kong, US, Switzerland: cross-listing codes', 'A-share codes on the Hong Kong table come from Eastmoney company profiles. Hong Kong ↔ US and Switzerland ↔ US links are matched by ISIN or English company name and only link the same legal entity.'],
 ['Source 1: A-share constituents', 'SWS Research, index constituents https://www.swsresearch.com/institute_sw/allIndex/releasedIndex/releasedetail?code=801150'],
 ['Source 2: quotes, Chinese names', 'Tencent Securities https://gu.qq.com/'],
 ['Source 3: A-share financials, Eastmoney industries', 'Eastmoney Data Center, earnings reports https://data.eastmoney.com/bbsj/202606/yjbb.html'],
 ['Source 4: company profiles', 'Eastmoney company profile https://emweb.securities.eastmoney.com/pc_hsf10/pages/index.html?type=web&code=SH688235#/gsgk'],
 ['Source 5: A+H list', 'Tencent Securities A+H shares https://stockapp.finance.qq.com/mstats/#mod=list&id=hk_ah'],
 ['Source 6: US tickers', 'SEC company tickers https://www.sec.gov/files/company_tickers_exchange.json'],
 ['Source 7: Hong Kong, US, Switzerland', 'TradingView Stock Screener https://www.tradingview.com/screener/'],
 ['Source 8: US added stocks', 'Nasdaq Stock Screener (Health Care) https://www.nasdaq.com/market-activity/stocks/screener'],
]
data = json.dumps({'markets': markets, 'notes': notes}, ensure_ascii=False, separators=(',', ':')).replace('</', '<\\/')
html = open(os.path.join(HERE, 's2_pharma_stocks_template.html')).read().replace('/*__DATA__*/null', data)
open(OUT, 'w').write(html)
print('written', OUT, len(html.encode()) // 1024, 'KB')
