"""Append official websites without changing existing workbook data.

Usage: python3 scripts/add_s2_company_websites.py <workbook.xlsx>
Sources: Eastmoney CompanySurvey (A), CompanyProfile (HK), and TradingView
symbol profiles' web_site_url (US/SIX). No domains are inferred from names.
Responses, including failures, persist as JSON files in /tmp/s2-company-websites-cache/.
"""
import concurrent.futures
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
import threading
import time
from urllib.parse import urlsplit, urlunsplit

from openpyxl import load_workbook

CACHE = Path('/tmp/s2-company-websites-cache')
LOCK = threading.Lock()
DENIED = ('eastmoney.com', 'tradingview.com', 'yahoo.com', 'wikipedia.org',
          'facebook.com', 'linkedin.com', 'twitter.com', 'x.com', 'instagram.com',
          'youtube.com', 'hkex.com.hk', 'nasdaq.com', 'nyse.com', 'six-group.com')


def normalise(value):
    if not isinstance(value, str):
        return None
    value = value.strip()
    if not value or re.search(r'[\s<>";，；]', value):
        return None
    if not re.match(r'^https?://', value, re.I):
        value = 'https://' + value.removeprefix('//')
    try:
        p = urlsplit(value)
        host = (p.hostname or '').lower()
        if '.' not in host or p.username or p.password or p.port:
            return None
        if any(host == d or host.endswith('.' + d) for d in DENIED):
            return None
        return urlunsplit(('https', host, p.path, p.query, ''))
    except ValueError:
        return None


def fetch(url):
    key = hashlib.sha256(url.encode()).hexdigest()
    entry = CACHE / (key + '.json')
    if entry.exists():
        return json.loads(entry.read_text()).get('body', '')
    attempts = []
    for attempt in range(3):  # initial request plus at most two retries
        time.sleep(0.4 + attempt)
        p = subprocess.run(['curl', '-sS', '-L', '--max-time', '20',
                            '-A', 'Mozilla/5.0', '-w', '\n%{http_code}', url],
                           capture_output=True, text=True)
        body, _, status = p.stdout.rpartition('\n')
        attempts.append({'status': status, 'body': body, 'error': p.stderr})
        with LOCK:
            entry.write_text(json.dumps({'url': url, 'body': body, 'attempts': attempts.copy()}, ensure_ascii=False))
        if p.returncode == 0 and status == '200':
            return body
    return ''


def source(sheet, code, exchange):
    if sheet == 'A股':
        prefix = 'BJ' if code.startswith(('4', '8', '92')) else 'SH' if code.startswith(('6', '9')) else 'SZ'
        return f'https://emweb.securities.eastmoney.com/PC_HSF10/CompanySurvey/PageAjax?code={prefix}{code}'
    if sheet == '港股':
        return f'https://emweb.securities.eastmoney.com/PC_HKF10/CompanyProfile/PageAjax?code={code.split(".")[0].zfill(5)}'
    exchange = 'SIX' if sheet == '瑞士' else {'AMEX': 'AMEX', 'NYSE': 'NYSE', 'NASDAQ': 'NASDAQ'}[exchange]
    return f'https://www.tradingview.com/symbols/{exchange}-{code.replace(".", "-")}/'


def lookup(task):
    row, sheet, code, exchange = task
    url = source(sheet, code, exchange)
    body = fetch(url)
    try:
        if sheet == 'A股':
            value = json.loads(body)['jbzl'][0].get('ORG_WEB')
        elif sheet == '港股':
            value = json.loads(body)['gszl'].get('gswz')
        else:
            match = re.search(r'"web_site_url"\s*:\s*("(?:[^"\\]|\\.)*")', body)
            value = json.loads(match[1]) if match else None
        return row, normalise(value)
    except (ValueError, KeyError, IndexError, TypeError):
        return row, None


def main():
    path = Path(sys.argv[1]).expanduser()
    wb = load_workbook(path)
    CACHE.mkdir(exist_ok=True)
    for sheet in ('A股', '港股', '美股', '瑞士'):
        ws = wb[sheet]
        heads = [c.value for c in ws[1]]
        column = heads.index('公司网址') + 1 if '公司网址' in heads else len(heads) + 1
        if column != len(heads) + ('公司网址' not in heads):
            raise ValueError('公司网址 must be the last column')
        ws.cell(1, column, '公司网址')
        tasks = [(i, sheet, str(ws.cell(i, heads.index('代码') + 1).value),
                  ws.cell(i, heads.index('交易所') + 1).value if '交易所' in heads else None)
                 for i in range(2, ws.max_row + 1)]
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            for n, (row, value) in enumerate(pool.map(lookup, tasks), 1):
                ws.cell(row, column).value = value
                if n % 100 == 0:
                    print(f'{sheet}: {n}/{len(tasks)}', flush=True)
        empty = [code for row, _, code, _ in tasks if not ws.cell(row, column).value]
        print(f'{sheet}: {len(tasks)-len(empty)}/{len(tasks)} filled; empty: {empty}', flush=True)
    wb.save(path)


if __name__ == '__main__':
    main()
